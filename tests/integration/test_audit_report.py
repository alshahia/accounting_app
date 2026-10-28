"""Audit-trail report: JSON view and downloads (report 21 §6, analysis/20 §7)."""

from __future__ import annotations

import io
from collections.abc import AsyncIterator

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncEngine

from app.shared.ids import new_id
from tests.integration.conftest import (
    SeededIdentity,
    delete_identity,
    seed_identity,
)
from tests.integration.test_returns import _headers

pytestmark = pytest.mark.integration

AUDIT_URL = "/api/v1/reports/audit-trail"
WIDE_RANGE = {"start": "2000-01-01", "end": "2100-12-31"}


@pytest.fixture
async def auditor(engine: AsyncEngine) -> AsyncIterator[SeededIdentity]:
    """A tenant whose only user holds the auditor role (audit.view)."""
    slug = f"audit-{new_id().hex[-12:]}"
    seeded = await seed_identity(
        engine,
        tenant_slug=slug,
        email=f"auditor-{slug}@example.com",
        password="correct horse battery staple",
        role_code="auditor",
    )
    try:
        yield seeded
    finally:
        await delete_identity(engine, seeded.tenant_id)


async def test_audit_trail_report_lists_recent_events(
    client: AsyncClient, auditor: SeededIdentity
) -> None:
    """The report returns the tenant trail newest first with report metadata."""
    headers = await _headers(client, auditor, auditor.company_a)
    response = await client.get(AUDIT_URL, params=WIDE_RANGE, headers=headers)
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["meta"]["generated_by"] == auditor.email
    assert body["meta"]["filters"]["start"] == WIDE_RANGE["start"]
    assert body["total"] == len(body["rows"]) > 0
    assert body["start_date"] == WIDE_RANGE["start"]
    assert body["end_date"] == WIDE_RANGE["end"]
    actions = [row["action"] for row in body["rows"]]
    assert any(action.startswith("auth.") for action in actions)
    times = [row["created_at"] for row in body["rows"]]
    assert times == sorted(times, reverse=True)


async def test_audit_trail_report_filters_and_rejects_a_bad_range(
    client: AsyncClient, auditor: SeededIdentity
) -> None:
    """Action filters narrow the rows and an inverted range is rejected."""
    headers = await _headers(client, auditor, auditor.company_a)
    unfiltered = await client.get(AUDIT_URL, params=WIDE_RANGE, headers=headers)
    assert unfiltered.status_code == 200, unfiltered.text
    actions = {
        row["action"] for row in unfiltered.json()["rows"] if row["action"].startswith("auth.")
    }
    assert actions, "the login must leave an authentication audit event"
    action = sorted(actions)[0]

    response = await client.get(AUDIT_URL, params={**WIDE_RANGE, "action": action}, headers=headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total"] >= 1
    assert {row["action"] for row in body["rows"]} == {action}
    assert all(row["actor_user_id"] == str(auditor.user_id) for row in body["rows"])

    bad = await client.get(
        AUDIT_URL, params={"start": "2026-02-01", "end": "2026-01-01"}, headers=headers
    )
    assert bad.status_code == 422


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
async def test_audit_trail_report_downloads(
    client: AsyncClient, auditor: SeededIdentity, fmt: str
) -> None:
    """Both downloads render the same trail as a real PDF or right-to-left workbook."""
    headers = await _headers(client, auditor, auditor.company_a)
    response = await client.get(f"{AUDIT_URL}.{fmt}", params=WIDE_RANGE, headers=headers)
    assert response.status_code == 200, response.text
    if fmt == "pdf":
        assert response.content.startswith(b"%PDF")
        assert response.headers["content-type"].startswith("application/pdf")
    else:
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        workbook = load_workbook(filename=io.BytesIO(response.content))
        sheet = workbook.active
        assert sheet is not None
        assert sheet.sheet_view.rightToLeft is True
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        assert any(value == "الإجراء" for value in values)
        assert any(isinstance(value, str) and value.startswith("وقت الإنشاء: ") for value in values)
        assert any(isinstance(value, str) and value.startswith("المرشحات:") for value in values)
        assert sheet.max_row > 4
