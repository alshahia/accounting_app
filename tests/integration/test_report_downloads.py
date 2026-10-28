"""Arabic PDF and right-to-left Excel downloads for the report pack.

Every report has a JSON view and matching downloads; these tests prove the
downloads render from the same service query and produce a real PDF/Workbook.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncEngine

from tests.integration.conftest import SeededLedger, seed_ledger
from tests.integration.test_documents import MasterData, seed_master
from tests.integration.test_report_pack import _seed_documents
from tests.integration.test_returns import _headers

pytestmark = pytest.mark.integration

PDF_URLS = (
    ("/api/v1/reports/trial-balance.pdf", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/profit-and-loss.pdf", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/balance-sheet.pdf", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/ar-aging.pdf", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/ap-aging.pdf", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/inventory-valuation.pdf", {}),
    ("/api/v1/reports/stock-movement.pdf", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/tax-return.pdf", {"start": "2026-01-01", "end": "2026-01-31"}),
)
XLSX_URLS = (
    ("/api/v1/reports/trial-balance.xlsx", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/profit-and-loss.xlsx", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/balance-sheet.xlsx", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/ar-aging.xlsx", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/ap-aging.xlsx", {"as_of": "2026-01-31"}),
    ("/api/v1/reports/inventory-valuation.xlsx", {}),
    ("/api/v1/reports/stock-movement.xlsx", {"start": "2026-01-01", "end": "2026-01-31"}),
    ("/api/v1/reports/tax-return.xlsx", {"start": "2026-01-01", "end": "2026-01-31"}),
)


@pytest.fixture
async def ledger(engine: AsyncEngine, identity) -> SeededLedger:
    """The identity fixture plus a chart and periods."""
    return await seed_ledger(engine, identity)


@pytest.fixture
async def master(engine: AsyncEngine, ledger: SeededLedger) -> MasterData:
    """Customer, supplier, tax code and item for the ledger's company."""
    return await seed_master(engine, ledger.identity)


@pytest.mark.parametrize(("url", "params"), PDF_URLS)
async def test_pdf_downloads(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
    url: str,
    params: dict[str, str],
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    response = await client.get(url, params=params, headers=headers)
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"%PDF")
    assert "attachment" in response.headers["content-disposition"]


@pytest.mark.parametrize(("url", "params"), XLSX_URLS)
async def test_xlsx_downloads_are_right_to_left(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
    url: str,
    params: dict[str, str],
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    response = await client.get(url, params=params, headers=headers)
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    assert sheet.sheet_view.rightToLeft is True
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert any(isinstance(value, str) and value.startswith("وقت الإنشاء: ") for value in values)
    assert any(isinstance(value, str) and value.startswith("أنشئ بواسطة: ") for value in values)
    assert any(isinstance(value, str) and value.startswith("الشركة: ") for value in values)


async def test_statement_and_ledger_downloads(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    params = {"start": "2026-01-01", "end": "2026-01-31"}
    statement = await client.get(
        "/api/v1/reports/customer-statement.pdf",
        params={"customer_id": str(master.customer_id), **params},
        headers=headers,
    )
    assert statement.status_code == 200, statement.text
    assert statement.content.startswith(b"%PDF")
    statement_xlsx = await client.get(
        "/api/v1/reports/customer-statement.xlsx",
        params={"customer_id": str(master.customer_id), **params},
        headers=headers,
    )
    assert statement_xlsx.status_code == 200, statement_xlsx.text
    assert statement_xlsx.content.startswith(b"PK")

    ledger_pdf = await client.get(
        "/api/v1/reports/general-ledger.pdf",
        params={"account_id": str(ledger.account_ids["1100"]), **params},
        headers=headers,
    )
    assert ledger_pdf.status_code == 200, ledger_pdf.text
    assert ledger_pdf.content.startswith(b"%PDF")


async def test_statement_comparative_downloads(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
) -> None:
    """The statement downloads print the prior period above the requested one."""
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    params = {
        "customer_id": str(master.customer_id),
        "start": "2026-02-01",
        "end": "2026-02-28",
        "compare": "true",
    }
    statement = await client.get(
        "/api/v1/reports/customer-statement.pdf", params=params, headers=headers
    )
    assert statement.status_code == 200, statement.text
    assert statement.content.startswith(b"%PDF")
    assert "attachment" in statement.headers["content-disposition"]

    workbook_response = await client.get(
        "/api/v1/reports/customer-statement.xlsx", params=params, headers=headers
    )
    assert workbook_response.status_code == 200, workbook_response.text
    workbook = load_workbook(BytesIO(workbook_response.content))
    sheet = workbook.active
    assert sheet is not None
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert any(
        isinstance(value, str) and value.startswith("الفترة السابقة:") for value in values
    ), values
    assert any(
        isinstance(value, str) and value.startswith("الفترة الحالية:") for value in values
    ), values
    assert any(
        isinstance(value, str) and value.startswith("الرصيد الختامي السابق:") for value in values
    ), values
    assert any(
        isinstance(value, str) and value.startswith("المرشحات:") and "compare=true" in value
        for value in values
    ), values


async def test_balance_sheet_pdf_is_a_substantial_document(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    response = await client.get(
        "/api/v1/reports/balance-sheet.pdf",
        params={"as_of": date(2026, 1, 31).isoformat()},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert len(response.content) > 1000


COMPARE_DOWNLOADS = (
    (
        "/api/v1/reports/trial-balance.pdf",
        {"start": "2026-01-01", "end": "2026-01-31", "compare": "true"},
    ),
    (
        "/api/v1/reports/profit-and-loss.pdf",
        {"start": "2026-01-01", "end": "2026-01-31", "compare": "true"},
    ),
    ("/api/v1/reports/balance-sheet.pdf", {"as_of": "2026-01-31", "compare": "true"}),
)


@pytest.mark.parametrize(("url", "params"), COMPARE_DOWNLOADS)
async def test_comparative_pdf_downloads(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
    url: str,
    params: dict[str, str],
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    response = await client.get(url, params=params, headers=headers)
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"%PDF")
    assert "attachment" in response.headers["content-disposition"]


COMPARE_WORKBOOKS = (
    (
        "/api/v1/reports/trial-balance.xlsx",
        {"start": "2026-01-01", "end": "2026-01-31", "compare": "true"},
        "الرصيد السابق",
    ),
    (
        "/api/v1/reports/profit-and-loss.xlsx",
        {"start": "2026-01-01", "end": "2026-01-31", "compare": "true"},
        "المبلغ السابق",
    ),
    (
        "/api/v1/reports/balance-sheet.xlsx",
        {"as_of": "2026-01-31", "compare": "true"},
        "المبلغ السابق",
    ),
)


@pytest.mark.parametrize(("url", "params", "prior_header"), COMPARE_WORKBOOKS)
async def test_comparative_workbooks_carry_prior_columns(
    client: AsyncClient,
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
    url: str,
    params: dict[str, str],
    prior_header: str,
) -> None:
    """A comparative workbook prints the prior-period column beside the current one."""
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    await _seed_documents(client, headers, master)
    response = await client.get(url, params=params, headers=headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert prior_header in values
    assert any(
        isinstance(value, str) and value.startswith("المرشحات:") and "compare=true" in value
        for value in values
    ), values
