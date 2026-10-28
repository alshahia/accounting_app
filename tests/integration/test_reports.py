"""Trial balance, Arabic PDF and Excel report tests (phase 0, day 8).

A known pair of posted entries gives a fixed trial balance; the JSON, the PDF
and the workbook must all reflect it, drafts must stay out, and the workbook must
open as a right-to-left sheet with numeric amounts.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncEngine, async_sessionmaker

from app.modules.ledger.service import JournalEntryInput, JournalLineInput, create_journal_entry
from app.shared.unit_of_work import UnitOfWork
from tests.integration.conftest import SeededIdentity, SeededLedger, seed_ledger

pytestmark = pytest.mark.integration

LOGIN_URL = "/api/v1/auth/login"
REPORT_URL = "/api/v1/reports/trial-balance"
JANUARY = date(2026, 1, 15)


@pytest.fixture
async def ledger(engine: AsyncEngine, identity: SeededIdentity) -> SeededLedger:
    """The identity fixture plus a chart and January/February periods."""
    return await seed_ledger(engine, identity)


@pytest.fixture
def unit_of_work(engine: AsyncEngine) -> UnitOfWork:
    """A unit of work over the test database."""
    return UnitOfWork(async_sessionmaker(engine, expire_on_commit=False))


async def _headers(
    client: AsyncClient, identity: SeededIdentity, company_id: Any
) -> dict[str, str]:
    response = await client.post(
        LOGIN_URL, json={"email": identity.email, "password": identity.password}
    )
    assert response.status_code == 200, response.text
    return {
        "Authorization": f"Bearer {response.json()['access_token']}",
        "X-Company-Id": str(company_id),
    }


async def _post_entries(ledger: SeededLedger, unit_of_work: UnitOfWork) -> None:
    identity = ledger.identity
    accounts = ledger.account_ids
    await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=JournalEntryInput(
            company_id=identity.company_a,
            period_id=ledger.period_id,
            entry_date=JANUARY,
            memo="sale for cash",
            post=True,
            lines=(
                JournalLineInput(account_id=accounts["1000"], debit_minor=1000),
                JournalLineInput(account_id=accounts["4000"], credit_minor=1000),
            ),
        ),
    )
    await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=JournalEntryInput(
            company_id=identity.company_a,
            period_id=ledger.period_id,
            entry_date=JANUARY,
            memo="expense",
            post=True,
            lines=(
                JournalLineInput(account_id=accounts["5900"], debit_minor=250),
                JournalLineInput(account_id=accounts["1000"], credit_minor=250),
            ),
        ),
    )
    await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=JournalEntryInput(
            company_id=identity.company_a,
            period_id=ledger.period_id,
            entry_date=JANUARY,
            memo="draft must not appear",
            post=False,
            lines=(
                JournalLineInput(account_id=accounts["1000"], debit_minor=9999),
                JournalLineInput(account_id=accounts["4000"], credit_minor=9999),
            ),
        ),
    )


async def test_trial_balance_json_is_balanced_and_excludes_drafts(
    client: AsyncClient,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> None:
    identity = ledger.identity
    await _post_entries(ledger, unit_of_work)
    headers = await _headers(client, identity, identity.company_a)
    response = await client.get(
        REPORT_URL, params={"start": "2026-01-01", "end": "2026-01-31"}, headers=headers
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["balanced"] is True
    assert body["total_debit_minor"] == body["total_credit_minor"] == 1250
    rows = {row["code"]: row for row in body["rows"]}
    assert rows["1000"]["debit_minor"] == 1000
    assert rows["1000"]["credit_minor"] == 250
    assert rows["1000"]["balance_minor"] == 750
    assert rows["4000"]["credit_minor"] == 1000
    assert rows["5900"]["debit_minor"] == 250
    meta = body["meta"]
    assert meta["generated_by"] == identity.email
    assert meta["company_id"] == str(identity.company_a)
    assert meta["company_name_ar"]
    assert meta["filters"] == {"start": "2026-01-01", "end": "2026-01-31"}
    assert meta["generated_at"]


async def test_trial_balance_pdf_is_a_real_pdf(
    client: AsyncClient,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> None:
    identity = ledger.identity
    await _post_entries(ledger, unit_of_work)
    headers = await _headers(client, identity, identity.company_a)
    response = await client.get(
        f"{REPORT_URL}.pdf",
        params={"start": "2026-01-01", "end": "2026-01-31"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert "trial-balance.pdf" in response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF")
    assert len(response.content) > 1500


async def test_trial_balance_excel_is_right_to_left_and_numeric(
    client: AsyncClient,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> None:
    identity = ledger.identity
    await _post_entries(ledger, unit_of_work)
    headers = await _headers(client, identity, identity.company_a)
    response = await client.get(
        f"{REPORT_URL}.xlsx",
        params={"start": "2026-01-01", "end": "2026-01-31"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    assert sheet.sheet_view.rightToLeft is True
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert any(isinstance(value, str) and value.startswith("أنشئ بواسطة: ") for value in values), (
        values
    )
    header_row = next(
        row for row in range(1, 20) if sheet.cell(row=row, column=1).value == "رمز الحساب"
    )
    assert header_row > 4, "metadata lines precede the header"
    assert sheet.cell(row=header_row, column=3).value == "مدين"
    by_code = {
        sheet.cell(row=row, column=1).value: row
        for row in range(header_row + 1, header_row + 11)
        if sheet.cell(row=row, column=1).value
    }
    cash_row = by_code["1000"]
    assert sheet.cell(row=cash_row, column=3).value == 10.0
    assert sheet.cell(row=cash_row, column=4).value == 2.5
    assert sheet.cell(row=cash_row, column=5).value == 7.5


async def test_report_range_validation_and_authentication(
    client: AsyncClient,
    ledger: SeededLedger,
) -> None:
    identity = ledger.identity
    headers = await _headers(client, identity, identity.company_a)
    invalid = await client.get(
        REPORT_URL, params={"start": "2026-02-01", "end": "2026-01-01"}, headers=headers
    )
    assert invalid.status_code == 422, invalid.text
    assert invalid.json()["code"] == "report_range_invalid"

    denied = await client.get(REPORT_URL, params={"start": "2026-01-01", "end": "2026-01-31"})
    assert denied.status_code == 401, denied.text
    assert denied.json()["code"] == "not_authenticated"


async def test_general_ledger_rows_carry_the_journal_entry_id_for_drill_down(
    client: AsyncClient,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> None:
    identity = ledger.identity
    await _post_entries(ledger, unit_of_work)
    headers = await _headers(client, identity, identity.company_a)
    response = await client.get(
        "/api/v1/reports/general-ledger",
        params={
            "account_id": str(ledger.account_ids["1000"]),
            "start": "2026-01-01",
            "end": "2026-01-31",
        },
        headers=headers,
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["meta"]["company_id"] == str(identity.company_a)
    assert body["meta"]["filters"]["account_id"] == str(ledger.account_ids["1000"])
    assert len(body["rows"]) == 2
    entry_ids = [row["entry_id"] for row in body["rows"]]
    assert all(entry_id for entry_id in entry_ids)
    assert len(set(entry_ids)) == 2
    for entry_id in entry_ids:
        detail = await client.get(f"/api/v1/ledger/journal-entries/{entry_id}", headers=headers)
        assert detail.status_code == 200, detail.text
        assert detail.json()["status"] == "posted"


async def test_trial_balance_compares_with_the_previous_period(
    client: AsyncClient,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> None:
    identity = ledger.identity
    await _post_entries(ledger, unit_of_work)
    headers = await _headers(client, identity, identity.company_a)
    response = await client.get(
        REPORT_URL,
        params={"start": "2026-02-01", "end": "2026-02-28", "compare": "true"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["start_date"] == "2026-02-01"
    assert body["compare_start_date"] == "2026-01-04"
    assert body["compare_end_date"] == "2026-01-31"
    assert body["compare_total_debit_minor"] == body["compare_total_credit_minor"] == 1250
    rows = {row["code"]: row for row in body["rows"]}
    assert rows["1000"]["debit_minor"] == 0
    assert rows["1000"]["compare_debit_minor"] == 1000
    assert rows["5900"]["compare_debit_minor"] == 250
    assert rows["4000"]["compare_credit_minor"] == 1000

    plain = await client.get(
        REPORT_URL,
        params={"start": "2026-02-01", "end": "2026-02-28"},
        headers=headers,
    )
    assert plain.status_code == 200, plain.text
    plain_body = plain.json()
    assert plain_body["compare_start_date"] is None
    assert plain_body["compare_total_debit_minor"] is None
    assert all(row["compare_debit_minor"] is None for row in plain_body["rows"])
