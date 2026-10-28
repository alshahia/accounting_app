"""Phase-0 acceptance run (day 10): A1-A13, A15 and the A16 demo path.

Run from the repository root against the development database:

    uv run python scripts/acceptance.py

It seeds temporary tenants, exercises the real application and services, runs
the automated guardrails, the resource measurement and the backup/restore
drill, and prints one PASS/FAIL line per acceptance criterion. A14 (Windows
virtual machine, offline install) cannot be executed from this host and is
recorded as unverified in the phase-0 exit review.

The script is self-cleaning: every tenant and database it creates is removed.
"""

from __future__ import annotations

import asyncio
import io
import subprocess
import sys
import time
import traceback
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import httpx  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pydantic import SecretStr  # noqa: E402
from sqlalchemy import func, select, text  # noqa: E402
from sqlalchemy.exc import DBAPIError  # noqa: E402
from sqlalchemy.ext.asyncio import (  # noqa: E402
    AsyncEngine,
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)
from sqlalchemy.pool import NullPool  # noqa: E402

from app.config import Settings  # noqa: E402
from app.main import create_app  # noqa: E402
from app.modules.inventory.models import StockMove  # noqa: E402
from app.modules.ledger.models import Account, JournalEntry, JournalLine  # noqa: E402
from app.modules.ledger.service import (  # noqa: E402
    JournalEntryInput,
    JournalLineInput,
    create_journal_entry,
)
from app.modules.purchasing.service import (  # noqa: E402
    PurchaseBillInput,
    PurchaseLineInput,
    create_purchase_bill,
)
from app.modules.reports.render import (  # noqa: E402
    render_trial_balance_pdf,
    render_trial_balance_xlsx,
)
from app.modules.reports.service import trial_balance  # noqa: E402
from app.modules.sales.service import (  # noqa: E402
    SalesInvoiceInput,
    SalesLineInput,
    create_sales_invoice,
)
from app.shared.errors import AppError  # noqa: E402
from app.shared.ids import new_id  # noqa: E402
from app.shared.unit_of_work import UnitOfWork  # noqa: E402
from tests.integration.conftest import (  # noqa: E402
    SeededIdentity,
    SeededLedger,
    delete_identity,
    seed_identity,
    seed_ledger,
)
from tests.integration.test_documents import (  # noqa: E402
    MasterData,
    _account_balances,
    _entry_lines,
    seed_master,
)

DATABASE_URL = "postgresql+asyncpg://postgres@127.0.0.1:55432/almahmoud"
JANUARY = date(2026, 1, 15)
FEBRUARY = date(2026, 2, 15)
BUDGET_SECONDS = 600


@dataclass
class Result:
    """One acceptance criterion outcome."""

    criterion: str
    title: str
    passed: bool
    evidence: str


RESULTS: list[Result] = []
FAILURES = 0


def record(criterion: str, title: str, passed: bool, evidence: str) -> None:
    global FAILURES
    RESULTS.append(Result(criterion, title, passed, evidence))
    if not passed:
        FAILURES += 1
    mark = "PASS" if passed else "FAIL"
    print(f"[{mark}] {criterion:3} {title}: {evidence}")


async def check(criterion: str, title: str, action: Any) -> None:
    try:
        evidence = await action
    except Exception as error:  # noqa: BLE001 - acceptance reports every failure
        record(criterion, title, False, f"{type(error).__name__}: {error}")
        traceback.print_exc()
    else:
        record(criterion, title, True, str(evidence))


def _code(exc: AppError) -> str:
    return f"{exc.status_code} {exc.code}"


async def expect_app_error(criterion: str, action: Any, expected_code: str) -> str:
    try:
        await action()
    except AppError as error:
        assert error.code == expected_code, f"expected {expected_code}, got {_code(error)}"
        return _code(error)
    raise AssertionError(f"expected AppError {expected_code}, nothing raised")


async def expect_db_error(action: Any) -> str:
    try:
        await action()
    except DBAPIError as error:
        message = str(error.orig).lower()
        return message.split("\n", 1)[0][:160]
    raise AssertionError("expected a database error, nothing raised")


async def _headers(client: httpx.AsyncClient, identity: SeededIdentity) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login",
        json={"email": identity.email, "password": identity.password},
    )
    assert response.status_code == 200, response.text
    return {
        "Authorization": f"Bearer {response.json()['access_token']}",
        "X-Company-Id": str(identity.company_a),
    }


async def check_a1(client: httpx.AsyncClient) -> str:
    health = await client.get("/api/v1/health")
    ready = await client.get("/api/v1/health/ready")
    assert health.status_code == 200, health.text
    assert ready.status_code == 200, ready.text
    body = health.json()
    return (
        f"{body['status']} v{body['version']} env={body['environment']}; "
        f"readiness {ready.json()}"
    )


async def check_a2(engine: AsyncEngine, tenant_a: SeededIdentity, tenant_b: SeededIdentity) -> str:
    async with engine.connect() as connection:
        transaction = await connection.begin()
        try:
            await connection.execute(text("SET ROLE almahmoud_app"))
            await connection.execute(
                text("SELECT set_config('app.tenant_id', :tenant, true)"),
                {"tenant": str(tenant_a.tenant_id)},
            )
            visible = (
                await connection.execute(
                    text("SELECT id FROM company WHERE id IN (:a, :b)"),
                    {"a": str(tenant_a.company_a), "b": str(tenant_b.company_a)},
                )
            ).scalars().all()
            assert visible == [tenant_a.company_a], f"cross-tenant read: {visible}"

            await connection.execute(
                text("SELECT set_config('app.tenant_id', '', true)")
            )
            none_visible = (
                await connection.execute(
                    text("SELECT id FROM company WHERE id IN (:a, :b)"),
                    {"a": str(tenant_a.company_a), "b": str(tenant_b.company_a)},
                )
            ).scalars().all()
            assert none_visible == [], f"fail-closed read leaked: {none_visible}"

            await connection.execute(
                text("SELECT set_config('app.tenant_id', :tenant, true)"),
                {"tenant": str(tenant_a.tenant_id)},
            )
            write_sql = text(
                "INSERT INTO company (id, tenant_id, code, name_ar, base_currency) "
                "VALUES (:id, :tenant, 'X-CROSS', 'cross', 'USD')"
            )
            try:
                await connection.execute(
                    write_sql,
                    {"id": str(new_id()), "tenant": str(tenant_b.tenant_id)},
                )
            except DBAPIError as error:
                write_denied = str(error.orig).lower()[:120]
            else:
                raise AssertionError("cross-tenant insert was accepted")
        finally:
            await transaction.rollback()
    return (
        f"tenant A read only A; unset tenant read 0 rows; cross-tenant insert refused "
        f"({write_denied})"
    )


async def check_a3(
    client: httpx.AsyncClient, identity: SeededIdentity, foreign: SeededIdentity
) -> str:
    anonymous = await client.get("/api/v1/audit/events")
    assert anonymous.status_code == 401, anonymous.text
    assert anonymous.json()["code"] == "not_authenticated"

    wrong = await client.post(
        "/api/v1/auth/login", json={"email": identity.email, "password": "wrong-password"}
    )
    assert wrong.status_code == 401, wrong.text

    headers = await _headers(client, identity)
    me = await client.get("/api/v1/auth/me", headers=headers)
    assert me.status_code == 200, me.text
    companies = me.json()["companies"]
    assert [company["id"] for company in companies] == [str(identity.company_a)]

    members = await client.get("/api/v1/admin/members", headers=headers)
    assert members.status_code == 403, members.text
    assert members.json()["code"] == "permission_denied"

    foreign_headers = {**headers, "X-Company-Id": str(foreign.company_a)}
    forbidden = await client.get("/api/v1/admin/members", headers=foreign_headers)
    assert forbidden.status_code == 403, forbidden.text
    assert forbidden.json()["code"] == "company_not_allowed"

    refresh = await client.post("/api/v1/auth/refresh")
    assert refresh.status_code == 403, refresh.text
    assert refresh.json()["code"] == "csrf_failed"
    return (
        "anonymous 401; bad password 401; /me 1 company; accountant members 403; "
        "foreign company 403 company_not_allowed; refresh without CSRF 403 csrf_failed"
    )


async def check_a4(
    engine: AsyncEngine,
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
) -> str:
    identity = ledger.identity
    accounts = ledger.account_ids

    async def unbalanced() -> None:
        await create_journal_entry(
            unit_of_work,
            tenant_id=identity.tenant_id,
            actor_user_id=identity.user_id,
            request=JournalEntryInput(
                company_id=identity.company_a,
                period_id=ledger.period_id,
                entry_date=JANUARY,
                post=True,
                lines=(
                    JournalLineInput(account_id=accounts["1000"], debit_minor=100),
                    JournalLineInput(account_id=accounts["4000"], credit_minor=90),
                ),
            ),
        )

    unbalanced_code = await expect_app_error("A4", unbalanced, "journal_unbalanced")

    async def closed_period() -> None:
        await create_journal_entry(
            unit_of_work,
            tenant_id=identity.tenant_id,
            actor_user_id=identity.user_id,
            request=JournalEntryInput(
                company_id=identity.company_a,
                period_id=ledger.closed_period_id,
                entry_date=FEBRUARY,
                post=True,
                lines=(
                    JournalLineInput(account_id=accounts["1000"], debit_minor=100),
                    JournalLineInput(account_id=accounts["4000"], credit_minor=100),
                ),
            ),
        )

    closed_code = await expect_app_error("A4", closed_period, "period_closed")

    source_id = new_id()

    def source_request() -> JournalEntryInput:
        return JournalEntryInput(
            company_id=identity.company_a,
            period_id=ledger.period_id,
            entry_date=JANUARY,
            post=True,
            source_type="acceptance_probe",
            source_id=source_id,
            lines=(
                JournalLineInput(account_id=accounts["1000"], debit_minor=50),
                JournalLineInput(account_id=accounts["3000"], credit_minor=50),
            ),
        )

    await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=source_request(),
    )

    async def duplicate_source() -> None:
        await create_journal_entry(
            unit_of_work,
            tenant_id=identity.tenant_id,
            actor_user_id=identity.user_id,
            request=source_request(),
        )

    source_code = await expect_app_error("A4", duplicate_source, "source_duplicate")

    async def raw_unbalanced() -> None:
        async with engine.begin() as connection:
            session = AsyncSession(bind=connection, expire_on_commit=False)
            entry = JournalEntry(
                tenant_id=identity.tenant_id,
                company_id=identity.company_a,
                period_id=ledger.period_id,
                number=f"ACC-{new_id().hex[:20]}",
                entry_date=JANUARY,
                status="posted",
                memo="raw unbalanced probe",
                posted_at=datetime.now(UTC),
                posted_by=identity.user_id,
            )
            session.add(entry)
            await session.flush()
            session.add_all(
                [
                    JournalLine(
                        tenant_id=identity.tenant_id,
                        company_id=identity.company_a,
                        entry_id=entry.id,
                        line_no=1,
                        account_id=accounts["1000"],
                        debit_minor=100,
                        credit_minor=0,
                    ),
                    JournalLine(
                        tenant_id=identity.tenant_id,
                        company_id=identity.company_a,
                        entry_id=entry.id,
                        line_no=2,
                        account_id=accounts["4000"],
                        debit_minor=0,
                        credit_minor=10,
                    ),
                ]
            )
            await session.flush()

    raw_message = await expect_db_error(raw_unbalanced)
    assert "balance" in raw_message or "unbalanced" in raw_message, raw_message
    return (
        f"service {unbalanced_code}; {closed_code}; {source_code}; raw insert refused ({raw_message})"
    )


async def check_a5(engine: AsyncEngine, entry_id: UUID, line_id: UUID) -> str:
    async def update_posted() -> None:
        async with engine.begin() as connection:
            await connection.execute(
                text("UPDATE journal_entry SET memo = 'tampered' WHERE id = :id"),
                {"id": str(entry_id)},
            )

    update_message = await expect_db_error(update_posted)
    assert "immutable" in update_message, update_message

    async def delete_entry() -> None:
        async with engine.begin() as connection:
            await connection.execute(
                text("DELETE FROM journal_entry WHERE id = :id"), {"id": str(entry_id)}
            )

    delete_message = await expect_db_error(delete_entry)
    assert "immutable" in delete_message, delete_message
    return f"posted update refused ({update_message}); posted delete refused ({delete_message})"


async def check_a6(
    ledger: SeededLedger,
    unit_of_work: UnitOfWork,
    engine: AsyncEngine,
) -> str:
    identity = ledger.identity
    memo = f"idempotency-{new_id().hex[-10:]}"
    request = JournalEntryInput(
        company_id=identity.company_a,
        period_id=ledger.period_id,
        entry_date=JANUARY,
        memo=memo,
        post=True,
        idempotency_key=f"acceptance-{memo}",
        lines=(
            JournalLineInput(account_id=ledger.account_ids["1000"], debit_minor=77),
            JournalLineInput(account_id=ledger.account_ids["4000"], credit_minor=77),
        ),
    )
    first = await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=request,
    )
    second = await create_journal_entry(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=request,
    )
    assert second.replayed is True
    assert second.entry.id == first.entry.id
    async with engine.connect() as connection:
        count = await connection.scalar(
            select(func.count()).select_from(JournalEntry).where(JournalEntry.memo == memo)
        )
    assert count == 1, count
    return f"replayed=True, same entry, stored rows={count}"


async def check_a7(ledger: SeededLedger, unit_of_work: UnitOfWork, engine: AsyncEngine) -> str:
    identity = ledger.identity
    started = time.perf_counter()

    async def post(index: int) -> UUID:
        result = await create_journal_entry(
            unit_of_work,
            tenant_id=identity.tenant_id,
            actor_user_id=identity.user_id,
            request=JournalEntryInput(
                company_id=identity.company_a,
                period_id=ledger.period_id,
                entry_date=JANUARY,
                memo=f"acceptance-parallel-{index}",
                post=True,
                idempotency_key=f"acceptance-parallel-{new_id().hex}",
                lines=(
                    JournalLineInput(
                        account_id=ledger.account_ids["1000"], debit_minor=10 + index
                    ),
                    JournalLineInput(
                        account_id=ledger.account_ids["3000"], credit_minor=10 + index
                    ),
                ),
            ),
        )
        return result.entry.id

    entries = await asyncio.gather(*(post(index) for index in range(20)))
    elapsed = time.perf_counter() - started
    async with engine.connect() as connection:
        distinct_numbers = await connection.scalar(
            select(func.count(func.distinct(JournalEntry.number))).where(
                JournalEntry.tenant_id == identity.tenant_id,
                JournalEntry.memo.like("acceptance-parallel-%"),
            )
        )
        debit, credit = (
            await connection.execute(
                select(
                    func.coalesce(func.sum(JournalLine.debit_minor), 0),
                    func.coalesce(func.sum(JournalLine.credit_minor), 0),
                )
                .select_from(JournalLine)
                .join(JournalEntry, JournalEntry.id == JournalLine.entry_id)
                .where(JournalEntry.memo.like("acceptance-parallel-%"))
            )
        ).one()
    assert len(entries) == 20
    assert distinct_numbers == 20, distinct_numbers
    assert int(debit) == int(credit), (debit, credit)
    return f"20 postings in {elapsed:.2f}s, 20 distinct numbers, debit==credit={debit}"


async def check_a8(
    engine: AsyncEngine,
    ledger: SeededLedger,
    master: MasterData,
    unit_of_work: UnitOfWork,
) -> dict[str, Any]:
    identity = ledger.identity
    bill = await create_purchase_bill(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=PurchaseBillInput(
            company_id=identity.company_a,
            supplier_id=master.supplier_id,
            bill_date=JANUARY,
            post=True,
            lines=(
                PurchaseLineInput(
                    description="ten widgets",
                    quantity=Decimal("10"),
                    unit_price_minor=1000,
                    item_id=master.item_id,
                ),
            ),
        ),
    )
    assert bill.bill.status == "posted", bill.bill.status

    first_sale = await create_sales_invoice(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=SalesInvoiceInput(
            company_id=identity.company_a,
            customer_id=master.customer_id,
            invoice_date=JANUARY,
            post=True,
            lines=(
                SalesLineInput(
                    description="four widgets",
                    quantity=Decimal("4"),
                    unit_price_minor=2000,
                    item_id=master.item_id,
                ),
            ),
        ),
    )
    second_sale = await create_sales_invoice(
        unit_of_work,
        tenant_id=identity.tenant_id,
        actor_user_id=identity.user_id,
        request=SalesInvoiceInput(
            company_id=identity.company_a,
            customer_id=master.customer_id,
            invoice_date=JANUARY,
            post=True,
            lines=(
                SalesLineInput(
                    description="six widgets",
                    quantity=Decimal("6"),
                    unit_price_minor=2000,
                    item_id=master.item_id,
                ),
            ),
        ),
    )
    async with engine.connect() as connection:
        item_stock = (
            await connection.execute(
                text(
                    "SELECT quantity_on_hand, stock_value_minor FROM item WHERE id = :id"
                ),
                {"id": str(master.item_id)},
            )
        ).one()
        balances = await _account_balances(
            connection, tenant_id=identity.tenant_id, company_id=identity.company_a
        )
        move_count = await connection.scalar(
            select(func.count()).select_from(StockMove).where(StockMove.item_id == master.item_id)
        )
    assert int(item_stock.quantity_on_hand) == 0
    assert int(item_stock.stock_value_minor) == 0
    assert int(move_count) == 3, move_count
    assert balances["1200"] == 0, balances
    assert balances["5100"] == 10000, balances
    assert balances["4000"] == -20000, balances
    assert balances["1100"] == 20000, balances
    assert balances["2000"] == -10000, balances
    assert second_sale.cogs_entry is not None
    return {
        "balances": balances,
        "stock": (int(item_stock.quantity_on_hand), int(item_stock.stock_value_minor)),
        "moves": int(move_count),
        "bill_total": bill.bill.total_minor,
        "sales_totals": (
            first_sale.invoice.total_minor,
            second_sale.invoice.total_minor,
        ),
        "entry_id": second_sale.journal_entry.entry.id if second_sale.journal_entry else None,
        "line_id": second_sale.journal_entry.lines[0].id if second_sale.journal_entry else None,
    }


async def check_a9(engine: AsyncEngine, ledger: SeededLedger) -> str:
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        balance = await trial_balance(
            session,
            tenant_id=ledger.identity.tenant_id,
            company_id=ledger.identity.company_a,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 1, 31),
        )
    assert balance.balanced
    started = time.perf_counter()
    pdf = render_trial_balance_pdf(
        balance, company_name_ar="شركة المحمود التجريبية", company_name_en="Al-Mahmoud Demo"
    )
    elapsed = time.perf_counter() - started
    assert pdf.startswith(b"%PDF"), pdf[:8]
    assert b"/Type /Page" in pdf
    assert len(pdf) > 1500
    return f"{len(pdf)} bytes in {elapsed:.2f}s, PDF header and page objects present"


async def check_a10(engine: AsyncEngine, ledger: SeededLedger) -> str:
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        balance = await trial_balance(
            session,
            tenant_id=ledger.identity.tenant_id,
            company_id=ledger.identity.company_a,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 1, 31),
        )
    workbook_bytes = render_trial_balance_xlsx(
        balance, company_name_ar="شركة المحمود التجريبية", company_name_en="Al-Mahmoud Demo"
    )
    workbook = load_workbook(io.BytesIO(workbook_bytes))
    sheet = workbook.active
    assert sheet is not None
    assert sheet.sheet_view.rightToLeft is True
    assert sheet.cell(row=4, column=1).value == "رمز الحساب"
    assert sheet.cell(row=4, column=3).value == "مدين"
    cash = None
    for row in range(5, 80):
        if sheet.cell(row=row, column=1).value == "1000":
            cash = row
            break
    assert cash is not None, "account 1000 missing from the workbook"
    assert sheet.cell(row=cash, column=3).value is not None
    return (
        f"{len(workbook_bytes)} bytes, rightToLeft=True, header row 4, account 1000 on row {cash}"
    )


def run_command(name: str, command: list[str]) -> tuple[bool, str]:
    started = time.perf_counter()
    result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, check=False)
    elapsed = time.perf_counter() - started
    if result.returncode != 0:
        return False, f"{name} exit {result.returncode}: {result.stdout[-400:]}{result.stderr[-400:]}"
    return True, f"{name} ok ({elapsed:.1f}s)"


async def check_a11() -> str:
    commands = [
        ("ruff-format", ["uv", "run", "ruff", "format", "--check", "."]),
        ("ruff-check", ["uv", "run", "ruff", "check", "."]),
        ("pyright", ["uv", "run", "pyright"]),
        (
            "analysis-check",
            ["uv", "run", "python", "work/scripts/check_analysis.py", "--root", ".", "--evidence", "--docs"],
        ),
        ("licence-gate", ["uv", "run", "python", "scripts/check_licences.py", "--quiet"]),
        ("dependency-audit", ["uv", "run", "pip-audit", "--progress-spinner", "off"]),
    ]
    evidence: list[str] = []
    for name, command in commands:
        passed, detail = await asyncio.to_thread(run_command, name, command)
        evidence.append(detail)
        if not passed:
            raise AssertionError("; ".join(evidence))
    return "; ".join(evidence)


def run_capture(command: list[str]) -> tuple[int, str, str]:
    result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, check=False)
    return result.returncode, result.stdout, result.stderr


async def check_a13() -> str:
    code, stdout, stderr = await asyncio.to_thread(
        run_capture, ["uv", "run", "python", "scripts/measure_resources.py"]
    )
    if code != 0:
        raise AssertionError(f"measure script exit {code}: {stdout[-500:]}{stderr[-500:]}")
    lines = [line.strip() for line in stdout.splitlines() if line.strip()]
    budget = next((line for line in lines if "memory-budget" in line), "")
    accounted = next((line for line in lines if "accounted resident" in line), "")
    return f"{accounted}; {budget}"


async def check_a15() -> str:
    code, stdout, stderr = await asyncio.to_thread(
        run_capture, ["uv", "run", "python", "scripts/backup_restore_drill.py"]
    )
    if code != 0:
        raise AssertionError(f"restore drill exit {code}: {stdout[-500:]}{stderr[-500:]}")
    restore = next((line.strip() for line in stdout.splitlines() if "restore-ok" in line), "")
    return restore


async def check_a16(client: httpx.AsyncClient) -> str:
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "demo@almahmoud.local", "password": "Al-Mahmoud-Demo-2026"},
    )
    assert login.status_code == 200, login.text
    token = login.json()["access_token"]
    me = await client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"})
    assert me.status_code == 200, me.text
    company = me.json()["companies"][0]
    headers = {"Authorization": f"Bearer {token}", "X-Company-Id": company["id"]}

    invoices = await client.get("/api/v1/sales/invoices", headers=headers)
    assert invoices.status_code == 200, invoices.text
    rows = invoices.json()
    assert len(rows) >= 3, len(rows)
    assert all(row["number"].startswith("SI-") for row in rows)

    report = await client.get(
        "/api/v1/reports/trial-balance",
        params={"start": "2026-01-01", "end": "2026-12-31"},
        headers=headers,
    )
    assert report.status_code == 200, report.text
    body = report.json()
    assert body["balanced"] is True
    pdf = await client.get(
        "/api/v1/reports/trial-balance.pdf",
        params={"start": "2026-01-01", "end": "2026-12-31"},
        headers=headers,
    )
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    return (
        f"demo login ok, {len(rows)} invoices, trial balance balanced "
        f"({body['total_debit_minor']} minor units), PDF {len(pdf.content)} bytes"
    )


async def main() -> int:
    started = time.perf_counter()
    engine = create_async_engine(DATABASE_URL, poolclass=NullPool)
    identity = await seed_identity(
        engine,
        tenant_slug=f"acceptance-{new_id().hex[-12:]}",
        email=f"acceptance-{new_id().hex[-12:]}@example.com",
        password="acceptance password 123",
    )
    foreign = await seed_identity(
        engine,
        tenant_slug=f"acceptance-foreign-{new_id().hex[-12:]}",
        email=f"acceptance-foreign-{new_id().hex[-12:]}@example.com",
        password="acceptance password 123",
    )
    business = await seed_identity(
        engine,
        tenant_slug=f"acceptance-business-{new_id().hex[-12:]}",
        email=f"acceptance-business-{new_id().hex[-12:]}@example.com",
        password="acceptance password 123",
    )
    try:
        ledger = await seed_ledger(engine, identity)
        business_ledger = await seed_ledger(engine, business)
        business_master = await seed_master(engine, business)
        unit_of_work = UnitOfWork(async_sessionmaker(engine, expire_on_commit=False))
        settings = Settings(
            environment="test",
            database_url=SecretStr(DATABASE_URL),
            log_level="WARNING",
            jwt_secret=SecretStr("acceptance-secret-key-for-the-phase-zero-run"),
            cookie_secure=False,
        )
        application = create_app(settings)
        async with application.router.lifespan_context(application):
            transport = httpx.ASGITransport(app=application)
            async with httpx.AsyncClient(
                transport=transport, base_url="http://acceptance"
            ) as client:
                await check("A1", "stack boots and readiness", check_a1(client))
                await check("A2", "tenant isolation", check_a2(engine, identity, foreign))
                await check("A3", "auth and permissions", check_a3(client, identity, foreign))
                await check("A16", "demo walkthrough path", check_a16(client))

        await check("A4", "ledger invariants", check_a4(engine, ledger, unit_of_work))
        await check("A6", "posting idempotency", check_a6(ledger, unit_of_work, engine))
        await check("A7", "concurrent postings", check_a7(ledger, unit_of_work, engine))
        business_evidence = await check_a8(
            engine, business_ledger, business_master, unit_of_work
        )
        print(f"       business evidence: {business_evidence}")
        record(
            "A8",
            "document posting and stock/GL reconciliation",
            True,
            str(business_evidence),
        )
        entry_id = business_evidence.get("entry_id")
        line_id = business_evidence.get("line_id")
        assert entry_id is not None and line_id is not None
        await check("A5", "posted ledger immutability", check_a5(engine, entry_id, line_id))
        await check("A9", "Arabic PDF report", check_a9(engine, business_ledger))
        await check("A10", "right-to-left Excel report", check_a10(engine, business_ledger))
    finally:
        await delete_identity(engine, identity.tenant_id)
        await delete_identity(engine, foreign.tenant_id)
        await delete_identity(engine, business.tenant_id)
        await engine.dispose()

    await check("A11", "guardrails (ruff, pyright, analysis, licences, audit)", check_a11())
    await check("A13", "memory budget", check_a13())
    await check("A15", "backup and restore drill", check_a15())

    elapsed = time.perf_counter() - started
    record(
        "A12",
        "acceptance run under ten minutes",
        elapsed < BUDGET_SECONDS,
        f"{elapsed:.1f}s (budget {BUDGET_SECONDS}s)",
    )

    print()
    passed = sum(1 for result in RESULTS if result.passed)
    print(f"acceptance: {passed}/{len(RESULTS)} checks passed, {FAILURES} failed, {elapsed:.1f}s")
    if FAILURES:
        print("failed criteria:", ", ".join(r.criterion for r in RESULTS if not r.passed))
        return 1
    print("phase-0 acceptance: all executed criteria passed (A14 unverified: Windows VM offline smoke)")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
